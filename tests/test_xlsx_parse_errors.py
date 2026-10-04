"""Malformed worksheet XML must use the existing invalid-file response."""
import zipfile

import openpyxl


def test_xlsx_late_xml_parse_error_returns_422_and_closes_resources(
    client, auth, temp_root, monkeypatch,
):
    from backend import files

    target = temp_root / "malformed-sheet.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.append(["ordinary cell", 7])
    workbook.save(target)
    workbook.close()
    with zipfile.ZipFile(target) as archive:
        entries = [(info, archive.read(info)) for info in archive.infolist()]
    with zipfile.ZipFile(target, "w") as archive:
        for info, body in entries:
            if info.filename == "xl/worksheets/sheet1.xml":
                assert b"</row>" in body
                body = body.replace(b"</row>", b"</mismatched-row>", 1)
            archive.writestr(info, body)

    opened = []
    loaded = []
    open_file = files._open_response_file
    load_workbook = openpyxl.load_workbook

    def capture_open(path):
        stream, info = open_file(path)
        opened.append(stream)
        return stream, info

    def capture_load(*args, **kwargs):
        book = load_workbook(*args, **kwargs)
        loaded.append(book)
        return book

    monkeypatch.setattr(files, "_open_response_file", capture_open)
    monkeypatch.setattr(openpyxl, "load_workbook", capture_load)
    # Observe the HTTP response that a real client receives on server errors.
    monkeypatch.setattr(client._transport, "raise_server_exceptions", False)
    response = client.get("/api/files/xlsx", params={"path": target.name}, headers=auth)
    assert len(loaded) == 1, "the genuine workbook load must complete before lazy XML parsing fails"
    assert len(opened) == 1 and opened[0].closed
    assert loaded[0]._archive.fp is None
    assert response.status_code == 422
    assert "failed to parse spreadsheet" in response.json()["detail"]
