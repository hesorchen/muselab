"""XLSX parsing must retain the file whose archive budgets were validated."""

import zipfile

from fastapi import HTTPException
import openpyxl
import pytest


@pytest.fixture
def opened_xlsx_files(app_module, monkeypatch):
    from backend import files

    opened = []
    open_file = files._open_response_file

    def capture_open(path):
        stream, info = open_file(path)
        opened.append(stream)
        return stream, info

    monkeypatch.setattr(files, "_open_response_file", capture_open)
    yield opened
    assert all(stream.closed for stream in opened)


def _write_workbook(path, marker):
    workbook = openpyxl.Workbook()
    workbook.active.title = "Data"
    workbook.active.append([marker, 7])
    workbook.save(path)
    workbook.close()


def test_xlsx_parser_uses_validated_open_file_after_atomic_replacement(
    client, auth, temp_root, monkeypatch, opened_xlsx_files,
):
    from backend import files

    target = temp_root / "changing.xlsx"
    replacement = temp_root / "replacement.xlsx"
    _write_workbook(target, "VALIDATED_ORIGINAL")
    _write_workbook(replacement, "UNVALIDATED_REPLACEMENT")
    opened = opened_xlsx_files
    validated = []
    parsed = []
    workbooks = []
    validate = files.validate_xlsx_archive
    load_workbook = openpyxl.load_workbook

    def validate_then_replace(source):
        validate(source)
        validated.append(source)
        replacement.replace(target)

    def parse(source, **kwargs):
        parsed.append(source)
        workbook = load_workbook(source, **kwargs)
        workbooks.append(workbook)
        return workbook

    monkeypatch.setattr(files, "validate_xlsx_archive", validate_then_replace)
    monkeypatch.setattr(openpyxl, "load_workbook", parse)
    response = client.get("/api/files/xlsx", params={"path": target.name}, headers=auth)
    assert response.status_code == 200
    assert response.json()["sheets"][0]["rows"] == [["VALIDATED_ORIGINAL", "7"]]
    assert len(opened) == 1
    assert validated == parsed == opened
    assert opened[0].closed
    assert workbooks[0]._archive.fp is None

    # The replacement really is a different, valid workbook, not mocked data.
    replacement_book = load_workbook(target, read_only=True, data_only=True)
    try:
        assert replacement_book.active["A1"].value == "UNVALIDATED_REPLACEMENT"
    finally:
        replacement_book.close()


@pytest.mark.parametrize("kind", ["corrupt_archive", "unsupported_workbook", "unsafe_budget"])
def test_xlsx_rejection_closes_owned_file(
    client, auth, temp_root, opened_xlsx_files, kind,
):
    target = temp_root / "rejected.xlsx"
    if kind == "corrupt_archive":
        target.write_bytes(b"not a ZIP archive")
    else:
        with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr(
                "xl/sharedStrings.xml",
                b"x" * (2 * 1024 * 1024) if kind == "unsafe_budget" else b"fixture",
            )
    response = client.get("/api/files/xlsx", params={"path": target.name}, headers=auth)
    assert response.status_code == 422
    assert len(opened_xlsx_files) == 1
    assert opened_xlsx_files[0].closed
    detail = response.json()["detail"]
    assert "compression ratio" in detail if kind == "unsafe_budget" else "failed to parse spreadsheet" in detail


def test_xlsx_iteration_failure_closes_workbook_and_owned_file(
    client, auth, temp_root, monkeypatch, opened_xlsx_files,
):
    from openpyxl.worksheet._read_only import ReadOnlyWorksheet

    target = temp_root / "iteration.xlsx"
    _write_workbook(target, "ORIGINAL")
    workbooks = []
    load_workbook = openpyxl.load_workbook
    failure = RuntimeError("controlled worksheet iteration failure")

    def parse(*args, **kwargs):
        workbook = load_workbook(*args, **kwargs)
        workbooks.append(workbook)
        return workbook

    def fail_iteration(*args, **kwargs):
        raise failure

    monkeypatch.setattr(openpyxl, "load_workbook", parse)
    monkeypatch.setattr(ReadOnlyWorksheet, "iter_rows", fail_iteration)
    with pytest.raises(RuntimeError) as caught:
        client.get("/api/files/xlsx", params={"path": target.name}, headers=auth)
    assert caught.value is failure
    assert len(opened_xlsx_files) == 1
    assert opened_xlsx_files[0].closed
    assert workbooks[0]._archive.fp is None


def test_xlsx_archive_validation_retains_caller_owned_binary_file(tmp_path):
    from backend.spreadsheet_safety import validate_xlsx_archive

    target = tmp_path / "ordinary.xlsx"
    _write_workbook(target, "CALLER_OWNED")
    with target.open("rb") as stream:
        validate_xlsx_archive(stream)
        assert not stream.closed
        with pytest.raises(HTTPException, match="entry budget"):
            validate_xlsx_archive(stream, max_entries=0)
        assert not stream.closed
        workbook = openpyxl.load_workbook(stream, read_only=True, data_only=True)
        try:
            assert workbook.active["A1"].value == "CALLER_OWNED"
        finally:
            workbook.close()
        assert not stream.closed


def test_xlsx_owned_file_preserves_worksheet_truncation_budgets(
    client, auth, temp_root, monkeypatch, opened_xlsx_files,
):
    from backend import files

    target = temp_root / "limited.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "First"
    for _ in range(4):
        workbook.active.append(["long cell", "second", "third"])
    workbook.create_sheet("Second").append(["other sheet"])
    workbook.save(target)
    workbook.close()
    for key, value in [("XLSX_MAX_SHEETS", 1), ("XLSX_MAX_ROWS", 2),
                       ("XLSX_MAX_COLS", 2), ("XLSX_CELL_MAX_CHARS", 4)]:
        monkeypatch.setattr(files, key, value)
    response = client.get("/api/files/xlsx", params={"path": target.name}, headers=auth)
    assert response.status_code == 200
    body = response.json()
    assert body["sheets_truncated"] is True
    assert body["limits"] == {"max_rows": 2, "max_cols": 2, "max_sheets": 1}
    assert body["sheets"] == [{
        "name": "First", "rows": [["long…", "seco…"], ["long…", "seco…"]],
        "rows_truncated": True, "cols_truncated": True,
    }]
    assert len(opened_xlsx_files) == 1 and opened_xlsx_files[0].closed
